"""Focused offline checks; these do not substitute for a real Pandoc/rendering run."""
import argparse
import copy
import importlib.util
from pathlib import Path
import tempfile
import unittest

from convert import (local_image, reject_external_vega, style_docx, transform,
                     validate_chart, verify_docx, PROFILES)


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.args = argparse.Namespace(mmdc='missing-mmdc-for-test', puppeteer_config=None)

    def tearDown(self):
        self.temp.cleanup()

    def test_local_unicode_and_percent_encoded_path(self):
        image = self.root / '中文 图片.png'
        image.write_bytes(b'fixture')
        self.assertEqual(local_image('中文%20图片.png', self.root), image.resolve())
        for path in ['https://example.test/a.png', 'data:image/png;base64,abc', 'missing.png']:
            with self.assertRaises((ValueError, FileNotFoundError)):
                local_image(path, self.root)

    def test_line_break_and_ordinary_code_survive(self):
        code = {'t': 'CodeBlock', 'c': [['', ['python'], []], 'print("<br>")']}
        ast = {'blocks': [copy.deepcopy(code), {'t': 'Para', 'c': [
            {'t': 'RawInline', 'c': ['html', '<br />']}]}]}
        transform(ast, self.root, self.root, self.args)
        self.assertEqual(ast['blocks'][0], code)
        self.assertEqual(ast['blocks'][1]['c'][0], {'t': 'LineBreak'})

    def test_raw_html_and_missing_mermaid_fail(self):
        for node in [{'t': 'RawBlock', 'c': ['html', '<table></table>']},
                     {'t': 'CodeBlock', 'c': [['', ['mermaid'], []], 'graph TD; A-->B']}]:
            with self.assertRaises((ValueError, FileNotFoundError)):
                transform({'blocks': [node]}, self.root, self.root, self.args)

    def test_chart_validation(self):
        valid = {'type': 'bar', 'x': ['A', 'B'], 'series': [{'values': [1, 2]}]}
        validate_chart(valid)
        for bad in [dict(valid, type='unknown'), dict(valid, series=[{'values': [1]}]),
                    dict(valid, series=[{'values': [1, float('nan')]}]),
                    dict(valid, type='pie', series=[{'values': [-1, 2]}]),
                    dict(valid, type='scatter')]:
            with self.assertRaises(ValueError):
                validate_chart(bad)

    def test_vega_external_resources_rejected(self):
        reject_external_vega({'data': {'values': [{'x': 1}]}, 'mark': 'bar'})
        for bad in [{'data': {'url': 'https://example.test/data.csv'}},
                    {'layer': [{'encoding': {'url': {'value': 'file:///secret'}}}]}]:
            with self.assertRaises(ValueError):
                reject_external_vega(bad)

    @unittest.skipUnless(importlib.util.find_spec('PIL') and importlib.util.find_spec('docx'),
                         'Requires existing Pillow and python-docx')
    def test_image_conversion_and_real_docx_structure(self):
        from PIL import Image
        from docx import Document
        from docx.shared import Inches
        source = self.root / '原图.bmp'
        Image.new('RGB', (80, 40), 'white').save(source)
        ast = {'blocks': [{'t': 'Para', 'c': [{'t': 'Image', 'c': [
            ['', [], []], [], [source.name, '']]}]}]}
        stats = transform(ast, self.root, self.root, self.args)
        image_path = Path(ast['blocks'][0]['c'][0]['c'][2][0])
        self.assertEqual(stats['images'], 1)
        with Image.open(image_path) as converted:
            self.assertEqual(converted.format, 'PNG')
        doc = Document()
        doc.add_heading('中文需求', 1)
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '字段'
        table.cell(1, 0).text = '过程\n下一步'
        doc.add_picture(str(image_path), width=Inches(2))
        output = self.root / '中文报告.docx'
        doc.save(output)
        style_docx(output, None)
        report = verify_docx(output)
        self.assertEqual(report['tables'], 1)
        self.assertEqual(report['drawings'], 1)
        self.assertEqual(report['embedded_image_relationships'], 1)
        self.assertIn('PAGE', report['fields'])
        reopened = Document(output)
        self.assertEqual(reopened.paragraphs[0].text, '中文需求')
        self.assertEqual(reopened.tables[0].cell(1, 0).text, '过程\n下一步')

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'Requires existing python-docx')
    def test_profiles_preserve_body_and_user_selected_title(self):
        from docx import Document
        from docx.oxml.ns import qn
        for profile in PROFILES:
            with self.subTest(profile=profile):
                doc = Document()
                doc.add_heading('产品能力', 1)
                doc.add_paragraph('客户端与服务端协作，权限边界不变。')
                table = doc.add_table(rows=2, cols=2)
                table.cell(0, 0).text = '能力'
                table.cell(1, 0).text = '告警处置'
                output = self.root / f'{profile}.docx'
                doc.save(output)
                style_docx(output, None, profile, '用户指定的产品需求说明书')
                result = Document(output)
                self.assertEqual(result.paragraphs[0].text, '用户指定的产品需求说明书')
                self.assertIn('客户端与服务端协作，权限边界不变。', [p.text for p in result.paragraphs])
                self.assertEqual(result.tables[0].cell(1, 0).text, '告警处置')
                self.assertEqual(result.core_properties.title, '用户指定的产品需求说明书')
                self.assertEqual(result.sections[0].header.paragraphs[0].text, '用户指定的产品需求说明书')
                self.assertEqual(result.styles['Normal'].paragraph_format.line_spacing,
                                 PROFILES[profile]['spacing'])
                self.assertIsNotNone(result.tables[0].rows[0]._tr.find('.//' + qn('w:tblHeader')))
                self.assertIn('PAGE', verify_docx(output)['fields'])

    @unittest.skipUnless(importlib.util.find_spec('docx'), 'Requires existing python-docx')
    def test_template_style_is_not_replaced_by_profile(self):
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.styles['Normal'].font.size = Pt(14)
        output = self.root / 'template-output.docx'
        doc.save(output)
        style_docx(output, 'existing-template.docx')
        self.assertEqual(Document(output).styles['Normal'].font.size.pt, 14)
        with self.assertRaises(ValueError):
            style_docx(output, 'existing-template.docx', 'formal', '标题')
        with self.assertRaises(ValueError):
            style_docx(output, None, 'formal', '')


if __name__ == '__main__':
    unittest.main()
