"""Regression cases for Markdown AST -> DOCX body parity checks."""
import tempfile
import json
import shutil
import subprocess
import unittest
from pathlib import Path

from document_content_check import compare, source_blocks


def word(text):
    return {'t': 'Str', 'c': text}


def plain(text):
    return {'t': 'Para', 'c': [word(text)]}


def heading(text):
    return {'t': 'Header', 'c': [1, ['', [], []], [word(text)]]}


class DocumentContentComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def ast(self):
        # Pandoc tables nest content blocks; cells must be included in parity.
        return {'blocks': [
            heading('第一章 方案'),
            plain('包含中文字符的正文'),
            {'t': 'Table', 'c': [[plain('字段名')], [plain('字段说明')]]},
            heading('第二章 实施'),
            plain('最后的验收内容'),
        ]}

    def docx(self, *, reorder=False, duplicate=False, omit=False, table=True):
        from docx import Document
        doc = Document()
        sections = [('第一章 方案', '包含中文字符的正文'),
                    ('第二章 实施', '最后的验收内容')]
        if reorder:
            sections.reverse()
        for title, body in sections:
            doc.add_heading(title, 1)
            doc.add_paragraph(body if not (omit and '最后' in body) else '意外的其他内容')
            if duplicate and title == '第一章 方案':
                doc.add_heading(title, 1)
            if title == '第一章 方案' and table:
                t = doc.add_table(rows=2, cols=1)
                t.cell(0, 0).text = '字段名'
                t.cell(1, 0).text = '字段说明'
        output = self.root / '测试 结果.docx'
        doc.save(output)
        return output

    def test_complete_docx(self):
        result = compare(self.ast(), self.docx())
        self.assertTrue(result['passed'], result)
        self.assertEqual(result['matched_text_blocks'], 6)
        self.assertEqual(result['tables_expected'], 1)

    def test_missing_body(self):
        result = compare(self.ast(), self.docx(omit=True))
        self.assertFalse(result['passed'])
        self.assertTrue(result['missing_or_out_of_order'])

    def test_heading_duplicate(self):
        result = compare(self.ast(), self.docx(duplicate=True))
        self.assertFalse(result['passed'])
        self.assertTrue(result['heading_differences'])

    def test_heading_out_of_order(self):
        result = compare(self.ast(), self.docx(reorder=True))
        self.assertFalse(result['passed'])
        self.assertFalse(result['heading_order_ok'])

    def test_table_missing(self):
        result = compare(self.ast(), self.docx(table=False))
        self.assertFalse(result['passed'])
        self.assertTrue(any('table count' in x for x in result['errors']))

    def test_split_runs_and_nontext_limits(self):
        from docx import Document
        doc = Document()
        doc.add_heading('第一章 方案', 1)
        paragraph = doc.add_paragraph()
        paragraph.add_run('拆分')
        paragraph.add_run('文本仍然完整')
        file = self.root / 'runs.docx'
        doc.save(file)
        ast = {'blocks': [heading('第一章 方案'),
                          {'t': 'Para', 'c': [word('拆分'), word('文本仍然完整')]}]}
        self.assertTrue(compare(ast, file)['passed'])
        ast['blocks'].append({'t': 'Para', 'c': [{'t': 'Math', 'c': [
            {'t': 'InlineMath'}, 'x^2']} ]})
        self.assertGreater(compare(ast, file)['excluded_from_text_comparison']['math'], 0)


    @unittest.skipUnless(shutil.which('pandoc'), 'Pandoc is not installed')
    def test_real_pandoc_markdown_docx_content(self):
        markdown = self.root / '中文 源文档.md'
        markdown.write_text(
            '# 报告\n\n## 一 说明\n\n这是正文示例。\n\n'
            '- 第一项\n- 第二项\n\n'
            '| 字段 | 说明 |\n| --- | --- |\n| id | 设备编号 |\n\n'
            '## 二 结论\n\n内容结束。\n', encoding='utf-8')
        docx = self.root / '实际 转换.docx'
        reader = 'markdown+pipe_tables+task_lists+tex_math_dollars+footnotes+raw_html'
        ast = json.loads(subprocess.check_output(
            ['pandoc', str(markdown), '-f', reader, '-t', 'json'], text=True))
        subprocess.run(['pandoc', str(markdown), '-f', reader,
                        '-t', 'docx', '--standalone', '-o', str(docx)],
                       check=True, timeout=60)
        result = compare(ast, docx)
        self.assertTrue(result['passed'], result)
        self.assertEqual(result['tables_expected'], 1)


if __name__ == '__main__':
    unittest.main()
