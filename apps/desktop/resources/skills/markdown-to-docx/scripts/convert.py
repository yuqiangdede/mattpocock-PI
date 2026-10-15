#!/usr/bin/env python3
"""Markdown -> Pandoc AST -> rendered assets -> DOCX; no automatic installs."""
import argparse
from collections import Counter
import importlib.util
import json
import math
import os
import re
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import unquote, urlsplit
import zipfile
import xml.etree.ElementTree as ET

READER = 'markdown+pipe_tables+task_lists+tex_math_dollars+footnotes+raw_html'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}


def windows_batch_command(argv, comspec=None):
    """Build cmd.exe's own command line (not MSVC list2cmdline quoting).

    Batch launchers are needed for npm's Windows .cmd shims. Quote all paths
    and reject characters that cmd.exe can expand or interpret unexpectedly.
    """
    argv = [str(x) for x in argv]
    if not argv:
        raise ValueError('Missing batch launcher')
    if any(not arg or any(c in arg for c in '\r\n\0"%&|<>^!') for arg in argv):
        raise ValueError('Unsafe Windows batch argument; use paths without cmd metacharacters')
    comspec = comspec or os.environ.get('COMSPEC', 'cmd.exe')
    if any(c in comspec for c in '\r\n\0"'):
        raise ValueError('Unsafe COMSPEC path')
    return f'"{comspec}" /d /s /c "' + ' '.join(f'"{arg}"' for arg in argv) + '"'


def run(argv, text=None, cwd=None):
    argv = [str(x) for x in argv]
    kw = {}
    if os.name == 'nt' and Path(argv[0]).suffix.lower() in {'.cmd', '.bat'}:
        # Passing a list would apply MSVC quoting to cmd.exe's /c argument a
        # second time. Pass a complete command line with an explicit launcher.
        kw['executable'] = os.environ.get('COMSPEC', 'cmd.exe')
        argv = windows_batch_command(argv, kw['executable'])
    result = subprocess.run(argv, input=text, encoding='utf-8', errors='strict',
                            stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            cwd=cwd, timeout=180, check=False, **kw)
    if result.returncode:
        raise RuntimeError(f'{argv if isinstance(argv, str) else argv[0]} failed: {result.stderr[-6000:]}')
    return result.stdout


def executable(name):
    found = shutil.which(name)
    if found:
        return found
    path = Path(name)
    if path.is_file():
        return str(path.resolve())
    raise FileNotFoundError(f'Missing executable: {name}; install only with authorization')


def nodes(value):
    if isinstance(value, dict):
        if 't' in value:
            yield value
        for child in value.values():
            yield from nodes(child)
    elif isinstance(value, list):
        for child in value:
            yield from nodes(child)


def local_image(target, base):
    # Windows drive letters are not URL schemes here.
    if not (len(target) > 2 and target[1] == ':'):
        parts = urlsplit(target)
        if parts.scheme or parts.netloc:
            raise ValueError(f'Localize remote/data image before conversion: {target}')
    path = Path(unquote(target))
    path = path if path.is_absolute() else base / path
    if not path.is_file():
        raise FileNotFoundError(f'Image not found: {path}')
    return path.resolve()


def validate_chart(spec):
    kind = spec.get('type')
    if kind not in {'bar', 'line', 'pie', 'scatter'}:
        raise ValueError('chart.type must be bar, line, pie or scatter')
    x, series = spec.get('x', []), spec.get('series', [])
    if not isinstance(x, list) or not x or not isinstance(series, list) or not series:
        raise ValueError('chart needs nonempty x and series arrays')
    for item in series:
        values = item.get('values', [])
        if len(values) != len(x) or not all(
                isinstance(y, (int, float)) and not isinstance(y, bool) and math.isfinite(y)
                for y in values):
            raise ValueError('chart series must contain finite numbers matching x length')
    if kind == 'pie' and (len(series) != 1 or min(series[0]['values']) < 0
                          or sum(series[0]['values']) <= 0):
        raise ValueError('pie needs one nonnegative series with positive total')
    if kind == 'scatter' and not all(isinstance(v, (int, float)) and
                                   not isinstance(v, bool) and math.isfinite(v) for v in x):
        raise ValueError('scatter x values must be finite numbers')


def render_chart(spec, output):
    validate_chart(spec)
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    if spec.get('font'):
        plt.rcParams['font.family'] = spec['font']
    fig, ax = plt.subplots(figsize=(8, 4.5), layout='constrained')
    try:
        x, series, kind = spec['x'], spec['series'], spec['type']
        for i, item in enumerate(series):
            y, label = item['values'], item.get('name', '')
            if kind == 'bar':
                width = .8 / len(series)
                ax.bar([j - .4 + width / 2 + i * width for j in range(len(x))],
                       y, width, label=label)
                ax.set_xticks(range(len(x)), x)
            elif kind == 'line':
                ax.plot(range(len(x)), y, marker='o', label=label)
                ax.set_xticks(range(len(x)), x)
            elif kind == 'scatter':
                ax.scatter(x, y, label=label)
            else:
                ax.pie(y, labels=x, autopct='%1.1f%%')
        ax.set(title=spec.get('title', ''), xlabel=spec.get('xlabel', ''),
               ylabel=spec.get('ylabel', ''))
        if kind != 'pie' and any(s.get('name') for s in series):
            ax.legend()
        fig.savefig(output, dpi=240)
    finally:
        plt.close(fig)


def reject_external_vega(value):
    if isinstance(value, dict):
        for key, item in value.items():
            if key in {'url', 'href', 'image'} or (key == 'data' and
                    (not isinstance(item, dict) or 'values' not in item)):
                raise ValueError('vega-lite must be self-contained with inline data.values')
            reject_external_vega(item)
    elif isinstance(value, list):
        for child in value:
            reject_external_vega(child)


def mermaid_render_flags(help_text, puppeteer_config=None):
    """Select options actually advertised by this mmdc installation."""
    options = set()
    for line in help_text.splitlines():
        # Only parse an option declaration, not option names in descriptions.
        spec = re.split(r'\s{2,}', line.strip(), maxsplit=1)[0]
        if spec.startswith('-'):
            options.update(re.findall(
                r'(?<![\w-])--[A-Za-z][\w-]*|(?<![\w-])-[A-Za-z](?![\w-])', spec))
    if not ({'-i', '--input'} & options) or not ({'-o', '--output'} & options):
        raise RuntimeError('Mermaid CLI help does not declare input/output flags')
    input_flag = '--input' if '--input' in options else '-i'
    output_flag = '--output' if '--output' in options else '-o'
    if '--size' in options:
        flags = ['--size', '1600']
    elif '--width' in options:
        flags = ['--width', '1600']
    elif '-w' in options:
        flags = ['-w', '1600']
    else:
        raise RuntimeError('Mermaid CLI does not advertise --size or --width')
    if '--scale' in options:
        flags.extend(['--scale', '2'])
    elif '-s' in options:
        flags.extend(['-s', '2'])
    if '--backgroundColor' in options:
        flags.extend(['--backgroundColor', 'white'])
    elif '-b' in options:
        flags.extend(['-b', 'white'])
    if puppeteer_config:
        if '--puppeteerConfigFile' in options:
            flags.extend(['--puppeteerConfigFile', str(Path(puppeteer_config).resolve())])
        elif '-p' in options:
            flags.extend(['-p', str(Path(puppeteer_config).resolve())])
        else:
            raise RuntimeError('Mermaid CLI does not advertise Puppeteer config option')
    return input_flag, output_flag, flags


def transform(ast, base, work, args):
    counter = Counter()
    mermaid_cli = None
    mermaid_flags = None
    input_flag = output_flag = None
    for node in nodes(ast):
        kind, content = node['t'], node.get('c')
        if kind == 'CodeBlock':
            classes = content[0][1]
            language = classes[0].lower() if classes else ''
            if language not in {'mermaid', 'chart', 'vega-lite'}:
                continue
            counter[language] += 1
            output = work / f'{language}-{counter[language]}.png'
            if language == 'mermaid':
                source = output.with_suffix('.mmd')
                source.write_text(content[1], encoding='utf-8')
                if mermaid_flags is None:
                    mermaid_cli = executable(args.mmdc)
                    input_flag, output_flag, mermaid_flags = mermaid_render_flags(
                        run([mermaid_cli, '--help']), args.puppeteer_config)
                run([mermaid_cli, input_flag, source, output_flag, output, *mermaid_flags])
            elif language == 'chart':
                render_chart(json.loads(content[1]), output)
            else:
                import vl_convert
                spec = json.loads(content[1])
                reject_external_vega(spec)
                output.write_bytes(vl_convert.vegalite_to_png(spec, scale=2))
            if not output.is_file() or output.stat().st_size == 0:
                raise RuntimeError(f'No rendered image produced: {output}')
            node.clear()
            node.update({'t': 'Para', 'c': [{'t': 'Image', 'c': [
                ['', [], []], [], [str(output), '']]}]})
        elif kind in {'RawInline', 'RawBlock'}:
            if kind == 'RawInline' and content[0] == 'html' and content[1].strip().lower() in {
                    '<br>', '<br/>', '<br />'}:
                node.clear()
                node['t'] = 'LineBreak'
            else:
                raise ValueError(f'Normalize unsupported raw markup first: {content[1][:100]}')
        elif kind == 'Image':
            path = local_image(content[2][0], base)
            suffix = path.suffix.lower()
            counter['images'] += 1
            if suffix not in {'.png', '.jpg', '.jpeg'}:
                output = work / f'image-{counter["images"]}.png'
                if suffix == '.svg':
                    import cairosvg
                    # Reject external SVG references before invoking a renderer.
                    svg = ET.parse(path)
                    for element in svg.iter():
                        for key, value in element.attrib.items():
                            if key.endswith('href') and not value.startswith('#'):
                                raise ValueError('SVG external references must be localized')
                        if element.tag.rsplit('}', 1)[-1] in {'style', 'script', 'foreignObject'}:
                            raise ValueError('Normalize SVG style/script/foreignObject before conversion')
                        if any('url(' in v.lower() for v in element.attrib.values()):
                            raise ValueError('Normalize SVG url() resources before conversion')
                    cairosvg.svg2png(url=str(path), write_to=str(output), scale=2)
                else:
                    from PIL import Image
                    with Image.open(path) as img:
                        img.seek(0)
                        img.convert('RGBA').save(output)
                    if suffix == '.gif':
                        counter['gif_first_frame'] += 1
                path = output
            content[2][0] = str(path)
    return dict(counter)


PROFILES = {
    'formal': {'size': 11, 'spacing': 1.25, 'margin': 2.5, 'accent': '244A73', 'cover': True},
    'review': {'size': 11, 'spacing': 1.5, 'margin': 2.8, 'accent': '1F4E79', 'cover': True},
    'compact': {'size': 10.5, 'spacing': 1.15, 'margin': 2.0, 'accent': '333333', 'cover': False},
}


def apply_profile(doc, profile, title):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    config = PROFILES[profile]
    if not title or not title.strip():
        raise ValueError('A user-specified --title is required for a profile')
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(config['margin'])
        section.left_margin = section.right_margin = Cm(config['margin'])
        header = section.header.paragraphs[0]
        header.text = title
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    normal = doc.styles['Normal']
    normal.font.size = Pt(config['size'])
    normal.paragraph_format.line_spacing = config['spacing']
    normal.paragraph_format.space_after = Pt(6 if profile != 'compact' else 3)
    for name in ['Normal', 'Title', 'Subtitle'] + [f'Heading {i}' for i in range(1, 7)]:
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        if name.startswith('Heading'):
            style.font.color.rgb = RGBColor.from_string(config['accent'])
            style.paragraph_format.keep_with_next = True
    for name in ['Source Code', 'Verbatim Char']:
        if name in doc.styles:
            doc.styles[name].font.name = 'Consolas'
    # Insert a title before the TOC/body without discarding original content.
    first = doc.element.body[0]
    heading = doc.add_paragraph(title, style='Title')
    first.addprevious(heading._p)
    heading.paragraph_format.space_before = Pt(120 if config['cover'] else 0)
    heading.paragraph_format.space_after = Pt(24)
    if config['cover']:
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        page_break = doc.add_page_break()
        heading._p.addnext(page_break._p)
    doc.core_properties.title = title
    for table in doc.tables:
        if 'Table Grid' in doc.styles:
            table.style = 'Table Grid'
        table_props = table._tbl.tblPr
        borders = table_props.find(qn('w:tblBorders'))
        if borders is None:
            borders = OxmlElement('w:tblBorders')
            table_props.append(borders)
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = borders.find(qn('w:' + edge))
            if border is None:
                border = OxmlElement('w:' + edge)
                borders.append(border)
            for key, value in {'val': 'single', 'sz': '4', 'color': 'CBD5E1'}.items():
                border.set(qn('w:' + key), value)
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                fill = (config['accent'] if profile == 'review' else 'E9EDF2') if row_index == 0 else None
                if fill:
                    props = cell._tc.get_or_add_tcPr()
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), fill)
                    props.append(shading)
                for paragraph in cell.paragraphs:
                    for run_item in paragraph.runs:
                        if row_index == 0:
                            run_item.bold = True
                            if profile == 'review':
                                run_item.font.color.rgb = RGBColor(255, 255, 255)


def style_docx(path, template, profile=None, title=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    doc = Document(path)
    if template and profile:
        raise ValueError('Use a reference template or a built-in profile, not both')
    if not template:
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.top_margin = section.bottom_margin = Cm(2.3)
            section.left_margin = section.right_margin = Cm(2.2)
        style = doc.styles['Normal']
        style.font.name, style.font.size = 'Microsoft YaHei', Pt(11)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        for section in doc.sections:
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), 'PAGE')
            section.footer.paragraphs[0]._p.append(field)
    if profile:
        apply_profile(doc, profile, title)
    # Nested tables are included, too.
    for table in doc.element.body.iter(qn('w:tbl')):
        row = table.find(qn('w:tr'))
        if row is not None:
            props = row.find(qn('w:trPr'))
            if props is None:
                props = OxmlElement('w:trPr')
                row.insert(0, props)
            if props.find(qn('w:tblHeader')) is None:
                props.append(OxmlElement('w:tblHeader'))
    settings = doc.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        settings.append(update)
    update.set(qn('w:val'), 'true')
    doc.save(path)


def verify_docx(path):
    with zipfile.ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError('Corrupted DOCX ZIP')
        xml_parts = []
        for name in archive.namelist():
            if name.endswith('.xml') or name.endswith('.rels'):
                parsed = ET.fromstring(archive.read(name))
                if name.startswith('word/') and name.endswith('.xml'):
                    xml_parts.append(parsed)
        document = ET.fromstring(archive.read('word/document.xml'))
        relationships = ET.fromstring(archive.read('word/_rels/document.xml.rels'))
        embedded = 0
        for rel in relationships:
            if rel.get('Type', '').endswith('/image'):
                if rel.get('TargetMode') == 'External':
                    raise ValueError('External image relationship in output')
                import posixpath
                target = rel.get('Target', '')
                name = target.lstrip('/') if target.startswith('/') else posixpath.normpath('word/' + target)
                if name not in archive.namelist():
                    raise ValueError(f'Missing embedded image: {target}')
                embedded += 1
        return {'paragraphs': len(document.findall('.//w:p', NS)),
                'tables': len(document.findall('.//w:tbl', NS)),
                'drawings': len(document.findall('.//a:blip', NS)),
                'embedded_image_relationships': embedded,
                'fields': [e.text for part in xml_parts for e in part.findall('.//w:instrText', NS)] +
                          [e.get('{'+NS['w']+'}instr') for part in xml_parts
                           for e in part.findall('.//w:fldSimple', NS)],
                'visual_check': 'not performed', 'field_update': 'requested, not verified'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', nargs='?', type=Path)
    parser.add_argument('--pandoc', default='pandoc')
    parser.add_argument('--mmdc', default='mmdc')
    parser.add_argument('--puppeteer-config')
    styling = parser.add_mutually_exclusive_group()
    styling.add_argument('--reference-doc', type=Path)
    styling.add_argument('--profile', choices=sorted(PROFILES))
    parser.add_argument('--title', help='Human-selected title used by a built-in profile')
    parser.add_argument('--toc', action='store_true')
    parser.add_argument('--overwrite', action='store_true')
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    if args.profile and (not args.title or not args.title.strip()):
        parser.error('--profile requires a user-specified --title')
    if args.title and not args.profile:
        parser.error('--title is used with --profile; edit the input title for template conversion')
    source = args.input.resolve()
    source.read_text(encoding='utf-8-sig')  # Validate readable UTF-8 before doing work.
    if args.check:
        result = {'pandoc': shutil.which(args.pandoc), 'mmdc': shutil.which(args.mmdc)}
        for module in ['docx', 'PIL', 'cairosvg', 'matplotlib', 'vl_convert']:
            result[module] = importlib.util.find_spec(module) is not None
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result['pandoc'] and result['docx'] else 2
    pandoc = executable(args.pandoc)
    output = (args.output or source.with_suffix('.docx')).resolve()
    if output.suffix.lower() != '.docx' or output == source:
        raise ValueError('Output must be a separate .docx file')
    if output.exists() and not args.overwrite:
        raise FileExistsError(f'Output exists; explicit --overwrite required: {output}')
    if args.reference_doc and not args.reference_doc.is_file():
        raise FileNotFoundError(args.reference_doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    work_root = output.parent / '.md-to-docx'
    work_root.mkdir(exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix=output.stem + '-', dir=work_root))
    print(f'Work directory: {work}', flush=True)
    ast = json.loads(run([pandoc, source, '-f', READER, '-t', 'json'], cwd=source.parent))
    counts = dict(Counter(n['t'] for n in nodes(ast)))
    render_counts = transform(ast, source.parent, work, args)
    ast_path = work / 'document.json'
    ast_path.write_text(json.dumps(ast, ensure_ascii=False), encoding='utf-8')
    temporary = work / 'document.docx'
    cmd = [pandoc, ast_path, '-f', 'json', '-t', 'docx', '-o', temporary, '--standalone']
    if args.toc or args.profile:
        cmd += ['--toc', '--toc-depth=3']
    if args.reference_doc:
        cmd += ['--reference-doc', args.reference_doc.resolve()]
    run(cmd, cwd=source.parent)
    style_docx(temporary, args.reference_doc, args.profile, args.title)
    from document_content_check import compare
    content_check = compare(ast, temporary)
    report = {'source': str(source), 'output': str(output), 'source_elements': counts,
              'profile': args.profile, 'title': args.title,
              'rendered': render_counts, 'docx': verify_docx(temporary),
              'content_comparison': content_check}
    if not content_check['passed']:
        raise ValueError(f'DOCX content mismatch: {content_check["errors"][:5]}; '
                         f'missing: {content_check["missing_or_out_of_order"][:3]}')
    expected = sum(n['t'] == 'Image' for n in nodes(ast))
    if report['docx']['drawings'] != expected:
        raise ValueError(f'Image count mismatch: AST {expected}, DOCX {report["docx"]["drawings"]}')
    report_path = work / 'report.json'
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    # All construction and structural checks succeeded before touching the target.
    if output.exists() and not args.overwrite:
        raise FileExistsError(output)
    os.replace(temporary, output)
    print(json.dumps({'output': str(output), 'report': str(report_path)}, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except (Exception,) as exc:
        print(f'Conversion failed: {exc}', file=sys.stderr)
        sys.exit(1)
